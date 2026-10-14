# -*- coding: utf-8 -*-
"""把《数字电路实验参考手册.md》构建为规范排版的 docx（交付版）。

用法: python build_docx.py
输出: 数字电路实验参考手册.docx（与本脚本同目录）

排版规范（参照高校实验指导书格式）:
  A4；页边距 上下 2.54cm / 左右 3.17cm；页脚居中页码
  一级标题 黑体 16pt 居中（另起一页）；二级 黑体 14pt；三级 黑体 12pt
  正文 宋体/Times New Roman 10.5pt（五号），1.5 倍行距
  代码 Consolas/宋体 9pt 浅灰底；表格 Table Grid 表头加粗浅蓝底
"""
import os
import re
import sys

from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

BASE = os.path.dirname(os.path.abspath(__file__))
SRC = os.path.join(BASE, "数字电路实验参考手册.md")
DST = os.path.join(BASE, "数字电路实验参考手册.docx")

EA = "宋体"
EA_HEAD = "黑体"
ASCII_FONT = "Times New Roman"
CODE_FONT = "Consolas"


# ---------------------------------------------------------------- 基础工具
def set_run(run, size=10.5, bold=False, italic=False, ascii_f=ASCII_FONT,
            ea_f=EA, color=None):
    run.font.name = ascii_f
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    if color:
        run.font.color.rgb = RGBColor(*color)
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = OxmlElement("w:rFonts")
        rPr.append(rFonts)
    rFonts.set(qn("w:ascii"), ascii_f)
    rFonts.set(qn("w:hAnsi"), ascii_f)
    rFonts.set(qn("w:eastAsia"), ea_f)


def shade_paragraph(p, hexcolor):
    pPr = p._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:fill"), hexcolor)
    pPr.append(shd)


def shade_cell(cell, hexcolor):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:fill"), hexcolor)
    tcPr.append(shd)


def add_page_break(doc):
    p = doc.add_paragraph()
    run = p.add_run()
    br = OxmlElement("w:br")
    br.set(qn("w:type"), "page")
    run._element.append(br)


def add_field(paragraph, instr, hint=""):
    r1 = paragraph.add_run()
    fld = OxmlElement("w:fldChar")
    fld.set(qn("w:fldCharType"), "begin")
    r1._element.append(fld)
    r2 = paragraph.add_run()
    it = OxmlElement("w:instrText")
    it.set(qn("xml:space"), "preserve")
    it.text = instr
    r2._element.append(it)
    r3 = paragraph.add_run()
    sep = OxmlElement("w:fldChar")
    sep.set(qn("w:fldCharType"), "separate")
    r3._element.append(sep)
    r4 = paragraph.add_run(hint)
    set_run(r4, 10.5, color=(128, 128, 128))
    r5 = paragraph.add_run()
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    r5._element.append(end)


INLINE_RE = re.compile(r"(\*\*.+?\*\*|`.+?`)")


def add_inline(p, text, size=10.5, base_bold=False):
    for tok in INLINE_RE.split(text):
        if not tok:
            continue
        if tok.startswith("**") and tok.endswith("**"):
            r = p.add_run(tok[2:-2])
            set_run(r, size, bold=True)
        elif tok.startswith("`") and tok.endswith("`"):
            r = p.add_run(tok[1:-1])
            set_run(r, size - 1, ascii_f=CODE_FONT, color=(0xB0, 0x30, 0x30))
        else:
            r = p.add_run(tok)
            set_run(r, size, bold=base_bold)


# ---------------------------------------------------------------- 文档骨架
def setup_page(doc):
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
    sec.top_margin = sec.bottom_margin = Cm(2.54)
    sec.left_margin = sec.right_margin = Cm(3.17)
    footer_p = sec.footer.paragraphs[0]
    footer_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_field(footer_p, " PAGE ", "")


def make_cover(doc):
    for _ in range(4):
        doc.add_paragraph()
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run(p.add_run("数字电路实验"), 26, bold=True, ea_f=EA_HEAD)
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run(p.add_run("实验参考手册"), 36, bold=True, ea_f=EA_HEAD)
    doc.add_paragraph()
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run(p.add_run("实验一　19 位并行前缀进位加法器"), 14, ea_f=EA_HEAD)
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run(p.add_run("实验二　8 位移位相加乘法器（Wallace 树）"), 14, ea_f=EA_HEAD)
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run(p.add_run("实验三　全流水脉动阵列矩阵乘法器"), 14, ea_f=EA_HEAD)
    for _ in range(6):
        doc.add_paragraph()
    for line in ("课程名称：数字电路实验", "适用对象：2026 秋 · 41 系",
                 "编制单位：课程教学组"):
        p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        set_run(p.add_run(line), 12)
    add_page_break(doc)


def make_toc(doc):
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run(p.add_run("目　　录"), 16, bold=True, ea_f=EA_HEAD)
    doc.add_paragraph()
    p = doc.add_paragraph()
    add_field(p, ' TOC \\o "1-2" \\h \\z \\u ',
              "（目录域：在 Word 中右键此处 → 更新域，即可自动生成目录）")
    add_page_break(doc)


# ---------------------------------------------------------------- md 解析渲染
def render_code_block(doc, lines):
    for idx, line in enumerate(lines):
        p = doc.add_paragraph()
        pf = p.paragraph_format
        pf.space_before = Pt(6 if idx == 0 else 0)
        pf.space_after = Pt(6 if idx == len(lines) - 1 else 0)
        pf.line_spacing = 1.0
        pf.left_indent = Cm(0.4)
        shade_paragraph(p, "F2F2F2")
        r = p.add_run(line if line else " ")
        set_run(r, 9, ascii_f=CODE_FONT)


def render_table(doc, rows):
    ncol = len(rows[0])
    table = doc.add_table(rows=len(rows), cols=ncol)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for ri, row in enumerate(rows):
        for ci in range(ncol):
            cell = table.cell(ri, ci)
            cell.text = ""
            p = cell.paragraphs[0]
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.space_after = Pt(1)
            header = ri == 0
            add_inline(p, row[ci], size=9.5, base_bold=header)
            if header:
                shade_cell(cell, "DCE6F1")
                for r in p.runs:
                    r.font.bold = True
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def render_image(doc, alt, path):
    full = os.path.join(BASE, path.replace("/", os.sep))
    if not os.path.exists(full):
        print("!! 图片缺失:", full)
        return
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    from PIL import Image
    with Image.open(full) as im:
        w_px = im.width
    width_cm = min(14.5, w_px / 150.0 * 2.54)   # 图按 150dpi 绘制
    run = p.add_run()
    run.add_picture(full, width=Cm(width_cm))
    cap = doc.add_paragraph(); cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.paragraph_format.space_after = Pt(10)
    set_run(cap.add_run(alt), 9, color=(0x59, 0x59, 0x59))


def build_body(doc, md_text):
    lines = md_text.splitlines()
    i, first_h1 = 0, True
    while i < len(lines):
        line = lines[i]
        s = line.strip()

        if s.startswith("```"):
            block = []
            i += 1
            while i < len(lines) and not lines[i].strip().startswith("```"):
                block.append(lines[i].rstrip())
                i += 1
            render_code_block(doc, block)
            i += 1
            continue

        if s.startswith("|"):
            rows = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                cells = [c.strip() for c in lines[i].strip().strip("|").split("|")]
                if not all(re.fullmatch(r":?-{2,}:?", c or "---") for c in cells):
                    rows.append(cells)
                i += 1
            if rows:
                render_table(doc, rows)
            continue

        m = re.match(r"^!\[(.*?)\]\((.*?)\)$", s)
        if m:
            render_image(doc, m.group(1), m.group(2))
            i += 1
            continue

        if s.startswith("# "):
            if not first_h1:
                add_page_break(doc)
            first_h1 = False
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            pf = p.paragraph_format
            pf.space_before, pf.space_after = Pt(6), Pt(18)
            p.style = doc.styles["Heading 1"]
            set_run(p.add_run(s[2:]), 16, bold=True, ea_f=EA_HEAD, color=(0, 0, 0))
            i += 1
            continue
        if s.startswith("## "):
            p = doc.add_paragraph()
            p.style = doc.styles["Heading 2"]
            p.paragraph_format.space_before = Pt(14)
            p.paragraph_format.space_after = Pt(6)
            set_run(p.add_run(s[3:]), 14, bold=True, ea_f=EA_HEAD, color=(0, 0, 0))
            i += 1
            continue
        if s.startswith("### "):
            p = doc.add_paragraph()
            p.style = doc.styles["Heading 3"]
            p.paragraph_format.space_before = Pt(10)
            p.paragraph_format.space_after = Pt(4)
            set_run(p.add_run(s[4:]), 12, bold=True, ea_f=EA_HEAD, color=(0, 0, 0))
            i += 1
            continue
        if s.startswith("> "):
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(0.75)
            p.paragraph_format.space_before = Pt(4)
            p.paragraph_format.space_after = Pt(4)
            pPr = p._p.get_or_add_pPr()
            pBdr = OxmlElement("w:pBdr")
            left = OxmlElement("w:left")
            left.set(qn("w:val"), "single"); left.set(qn("w:sz"), "18")
            left.set(qn("w:space"), "4"); left.set(qn("w:color"), "2E74B5")
            pBdr.append(left); pPr.append(pBdr)
            add_inline(p, s[2:], size=10)
            i += 1
            continue
        if s.startswith("- "):
            p = doc.add_paragraph(style="List Bullet")
            p.paragraph_format.space_after = Pt(2)
            add_inline(p, s[2:])
            i += 1
            continue
        if s == "---" or s == "":
            i += 1
            continue

        p = doc.add_paragraph()
        pf = p.paragraph_format
        pf.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
        pf.space_after = Pt(3)
        pf.first_line_indent = Cm(0.74)      # 首行缩进两字符
        add_inline(p, s)
        i += 1


def main():
    with open(SRC, encoding="utf-8") as f:
        md = f.read()
    doc = Document()
    st = doc.styles["Normal"]
    st.font.name = ASCII_FONT
    st.font.size = Pt(10.5)
    st.element.rPr.rFonts.set(qn("w:eastAsia"), EA)
    setup_page(doc)
    make_cover(doc)
    make_toc(doc)
    build_body(doc, md)
    props = doc.core_properties
    props.title = "数字电路实验参考手册"
    props.author = "课程教学组"
    doc.save(DST)
    print("已生成:", DST, "大小(KB):", os.path.getsize(DST) // 1024)


if __name__ == "__main__":
    sys.exit(main())